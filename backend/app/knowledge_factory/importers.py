"""Bulk import — turn heterogeneous sources into **draft** Knowledge Cards.

Supports Markdown, CSV, JSON, HTML and plain text natively; PDF / DOCX / XLSX via
optional libraries when present (a missing library degrades to a clear
"unsupported" result, never a crash — the project's "always usable" rule).

Every output is a Brouillon built on the official scaffold, with the imported
content parked in an ``## Import (à valider)`` section for an editor. **Nothing is
ever published**, and this module writes no files itself — it returns drafts.
"""

import csv
import io
import json
import re
from dataclasses import dataclass
from pathlib import Path

from app.knowledge_factory.generator import generate_card

_TAG_RE = re.compile(r"<[^>]+>")
_SLUG_RE = re.compile(r"[^a-z0-9]+")


@dataclass(frozen=True)
class DraftCard:
    slug: str
    text: str
    source: str


def slugify(value: str) -> str:
    slug = _SLUG_RE.sub("-", value.strip().lower()).strip("-")
    return slug or "sans-titre"


def _draft(*, title: str, profession: str, body: str, source: str) -> DraftCard:
    slug = slugify(title)
    scaffold = generate_card(slug=slug, title=title or slug, profession=profession or "a-preciser")
    text = f"{scaffold}\n## Import (à valider)\n\n{body.strip()}\n"
    return DraftCard(slug=slug, text=text, source=source)


def _rows_to_drafts(rows: list[dict], source: str) -> list[DraftCard]:
    drafts: list[DraftCard] = []
    for row in rows:
        lowered = {str(k).lower(): ("" if v is None else str(v)) for k, v in row.items()}
        title = lowered.get("title") or lowered.get("titre") or lowered.get("designation") or lowered.get("nom") or ""
        profession = lowered.get("profession") or lowered.get("metier") or lowered.get("trade") or ""
        body = lowered.get("body") or lowered.get("description") or lowered.get("resume") or json.dumps(row, ensure_ascii=False)
        if not title:
            continue
        drafts.append(_draft(title=title, profession=profession, body=body, source=source))
    return drafts


def import_bytes(name: str, data: bytes) -> tuple[list[DraftCard], str | None]:
    """Import one source by name+content. Returns ``(drafts, unsupported_reason)``;
    ``unsupported_reason`` is set (and drafts empty) when the format needs an
    optional library that is not installed."""
    suffix = Path(name).suffix.lower()
    text = data.decode("utf-8", errors="replace")

    if suffix in (".md", ".markdown", ".txt"):
        title = _first_heading(text) or Path(name).stem
        return [_draft(title=title, profession="", body=text, source=name)], None

    if suffix in (".html", ".htm"):
        stripped = _TAG_RE.sub(" ", text)
        title = _first_heading(text) or Path(name).stem
        return [_draft(title=title, profession="", body=stripped, source=name)], None

    if suffix == ".csv":
        rows = list(csv.DictReader(io.StringIO(text)))
        return _rows_to_drafts(rows, name), None

    if suffix == ".json":
        payload = json.loads(text)
        rows = payload if isinstance(payload, list) else [payload]
        rows = [r for r in rows if isinstance(r, dict)]
        return _rows_to_drafts(rows, name), None

    if suffix == ".pdf":
        return _import_pdf(name, data)
    if suffix == ".docx":
        return _import_docx(name, data)
    if suffix in (".xlsx", ".xls"):
        return _import_xlsx(name, data)

    return [], f"Format non supporté : {suffix or '∅'}"


def _first_heading(text: str) -> str:
    for line in text.splitlines():
        if line.startswith("# "):
            return line[2:].strip()
    return ""


def _import_pdf(name: str, data: bytes) -> tuple[list[DraftCard], str | None]:
    try:
        from pypdf import PdfReader  # type: ignore[import-not-found]
    except Exception:  # noqa: BLE001 — optional dependency absent
        return [], "Import PDF indisponible (installer 'pypdf')."
    reader = PdfReader(io.BytesIO(data))
    body = "\n".join((page.extract_text() or "") for page in reader.pages)
    return [_draft(title=Path(name).stem, profession="", body=body, source=name)], None


def _import_docx(name: str, data: bytes) -> tuple[list[DraftCard], str | None]:
    try:
        import docx  # type: ignore[import-not-found]
    except Exception:  # noqa: BLE001
        return [], "Import DOCX indisponible (installer 'python-docx')."
    document = docx.Document(io.BytesIO(data))
    body = "\n".join(p.text for p in document.paragraphs)
    return [_draft(title=Path(name).stem, profession="", body=body, source=name)], None


def _import_xlsx(name: str, data: bytes) -> tuple[list[DraftCard], str | None]:
    try:
        from openpyxl import load_workbook  # type: ignore[import-not-found]
    except Exception:  # noqa: BLE001
        return [], "Import Excel indisponible (installer 'openpyxl')."
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    header = [str(h or "") for h in next(rows_iter, [])]
    rows = [dict(zip(header, values, strict=False)) for values in rows_iter]
    return _rows_to_drafts(rows, name), None


__all__ = ["DraftCard", "import_bytes", "slugify"]
